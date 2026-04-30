from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).with_name("蜜蜂虚拟飞行行为仿真与可视化系统V1.0_软件设计说明书.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="9E9E9E", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, align=None, font_size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(font_size)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)


def set_default_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.first_line_indent = Pt(0)


def setup_section(section):
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.0)


def add_header_footer(section, title):
    header = section.header.paragraphs[0]
    header.text = title
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)
    add_page_number(section.footer.paragraphs[0])


def p(doc, text="", align=None, first_line=True):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    if not first_line:
        para.paragraph_format.first_line_indent = Pt(0)
    para.add_run(text)
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style=None)
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.left_indent = Pt(18)
    para.paragraph_format.line_spacing = 1.3
    para.add_run("● ").bold = True
    para.add_run(text)
    return para


def add_placeholder(doc, idx, title):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_border(cell, "A6A6A6", "10")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(24)
    para.paragraph_format.space_after = Pt(24)
    run = para.add_run(f"[图{idx}：{title}]")
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(96, 96, 96)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    cap = caption.add_run(f"图 {idx} {title}")
    cap.font.name = "宋体"
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    cap.font.size = Pt(10)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "D9EAF7")
        set_cell_text(cell, h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_border(cell, "7F7F7F")
        if widths:
            cell.width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), align=WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.CENTER)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[i], "BFBFBF", "6")
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    setup_section(doc.sections[0])
    set_default_styles(doc)

    # Cover page
    p(doc, "蜜蜂虚拟飞行行为仿真与可视化系统 V1.0", WD_ALIGN_PARAGRAPH.CENTER, False).runs[0].font.size = Pt(16)
    doc.add_paragraph()
    for word in ["软", "件", "设", "计", "说", "明", "书"]:
        para = p(doc, word, WD_ALIGN_PARAGRAPH.CENTER, False)
        para.runs[0].font.name = "黑体"
        para.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        para.runs[0].font.size = Pt(24)
        para.paragraph_format.space_after = Pt(10)
    doc.add_paragraph()
    p(doc, "天津理工大学", WD_ALIGN_PARAGRAPH.CENTER, False).runs[0].font.size = Pt(14)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(section)
    add_header_footer(section, "蜜蜂虚拟飞行行为仿真与可视化系统 V1.0 软件设计说明书")

    doc.add_heading("一、引言", level=1)
    doc.add_heading("1. 编写目的", level=2)
    for text in [
        "在昆虫行为学、智能体仿真和虚拟生态教学展示中，蜜蜂飞行行为具有速度快、姿态变化频繁、受环境扰动影响明显等特点，直接依靠人工动画或静态演示难以持续、直观地表现其目标搜索、趋近、悬停、采集和路径循环等过程。传统关键帧动画需要逐帧设定蜜蜂的位置和姿态，难以体现飞行控制、环境扰动、目标响应和可视化反馈之间的动态关系。",
        "为解决上述问题，本软件基于 Unity 引擎开发“蜜蜂虚拟飞行行为仿真与可视化系统 V1.0”，通过 C# 脚本构建蜜蜂飞行控制、目标搜索、环境扰动、轨迹记录、可视化展示和演示交互等模块，使虚拟蜜蜂能够在三维场景中依据控制输入、目标状态和环境约束实时生成运动轨迹。系统可用于蜜蜂飞行行为的教学演示、科普展示、虚拟仿真验证及相关交互应用开发。",
        "本软件的开发目的是提供一套轻量化、可交互、可观察的蜜蜂飞行行为仿真工具，使研究者、教育工作者和开发者能够在虚拟三维环境中实时驱动蜜蜂运动、调节飞行参数、观察目标趋近与采集行为、查看力与轨迹可视化反馈，从而降低昆虫行为仿真与可视化展示的技术门槛，为教学演示、科普传播和智能体行为建模提供直观的软件支撑。",
    ]:
        p(doc, text)

    doc.add_heading("2. 软件简介", level=2)
    for text in [
        "蜜蜂虚拟飞行行为仿真与可视化系统 V1.0 是一款基于 Unity 引擎开发的三维虚拟飞行行为仿真软件。系统以虚拟蜜蜂为核心对象，围绕“飞行控制—环境扰动—目标趋近—轨迹可视化—状态演示”的流程，模拟蜜蜂在三维空间中的手动飞行、自动巡航、目标接近、悬停采集、路径循环和避障修正等行为。",
        "在核心功能上，系统的飞行控制模块负责蜜蜂的运动驱动、振翅扰动模拟、飞行姿态调节、地面规避、障碍规避、目标吸引力响应以及运行时力向量的可视化显示；目标行为模块负责目标点搜索、自动目标选择、目标悬停、采集触发、路线循环与稳定路线复现；轨迹记录模块负责实时采样蜜蜂位置并绘制飞行轨迹曲线；可视化控制模块统一管理调试射线、力向量和轨迹线的显示与隐藏；演示交互模块提供运行时状态面板、飞行模式切换、操作帮助面板和参数控制入口，便于演示、截图和录屏。",
        "软件包含主菜单和仿真演示两个主要场景。主菜单场景提供系统首页、开始体验和退出等导航入口；仿真演示场景承载蜜蜂飞行运动、目标点布置、相机跟随观察、状态信息显示和可视化调试等核心运行内容。用户可通过键盘和界面控件进行飞行操作、模式切换和视角控制，从而直观理解蜜蜂飞行行为的生成过程。",
    ]:
        p(doc, text)

    doc.add_heading("3. 软件运行环境", level=2)
    for text in [
        "硬件环境方面，本软件要求处理器为 Intel Core i5 或 AMD Ryzen 5 及以上性能级别，内存容量不低于 8GB，硬盘可用空间不少于 2GB，并配备支持 DirectX 11 的独立或集成显卡。以上配置能够保证三维场景的实时渲染、物理计算和轨迹可视化在流畅帧率下稳定运行。",
        "软件环境方面，本软件基于 Windows 10 或 Windows 11 操作系统开发与运行，使用 Unity 2022.3.61f1c1 作为三维引擎和运行时环境，采用 Visual Studio 2022 作为代码编辑与编译工具，编程语言为 C#，目标平台为 Windows 独立可执行程序。",
        "显示与输入方面，推荐使用 Full HD（1920×1080）及以上分辨率的显示设备，以键盘和鼠标作为主要输入设备。软件包含两个主要场景文件：MainMenu.unity 用于主菜单界面展示与场景导航，Demo.unity 用于飞行仿真、可视化演示与交互操作。",
    ]:
        p(doc, text)

    doc.add_heading("二、软件使用步骤", level=1)
    doc.add_heading("1. 截图清单", level=2)
    p(doc, "以下截图为文档建议配图。当前文档已按照图号预留占位，待软件运行截图准备完成后，可直接替换相应占位区域。")
    screenshots = [
        ["图1", "软件主菜单界面", "展示软件名称、开始探索和退出等首页入口。"],
        ["图2", "仿真飞行场景", "展示三维场景、蜜蜂对象、手动/自动飞行控制及目标点。"],
        ["图3", "目标交互与行为表现", "展示蜜蜂趋近目标、悬停采集及避障修正等行为效果。"],
        ["图4", "轨迹与力向量可视化", "展示飞行轨迹线、各类力向量和运行时可视化内容。"],
        ["图5", "演示控制与状态面板", "展示 HUD 状态面板、力学参数控制和帮助说明界面。"],
    ]
    add_table(doc, ["图号", "截图内容", "对应功能点"], screenshots, widths=[1.6, 4.2, 9.9])

    steps = [
        ("启动与主界面", "用户启动软件进入主菜单，展示系统名称与“开始探索”等导航入口。点按“开始探索”后加载 Demo 场景，进入三维飞行仿真环境。", "软件主菜单界面"),
        ("仿真飞行与控制", "在仿真场景中，用户可通过键盘（W/S/A/D/Space/Ctrl/Shift）手动控制蜜蜂飞行，或切换至自动巡航模式由 BeeTargetController 搜索目标并自动生成飞行路径。BeeSimulation 实时计算控制力、目标引导力、姿态和速度，驱动蜜蜂在场景中飞行。", "仿真飞行场景"),
        ("目标交互与行为表现", "蜜蜂接近目标时进入悬停采集状态，在目标上方小幅波动模拟访花行为。自动巡航中系统通过前向射线检测障碍物和地面高度，实时计算规避力实现绕障和近地修正。", "目标交互与行为表现"),
        ("轨迹与力向量可视化", "用户可开启可视化显示：系统按固定间隔采样蜜蜂位置并绘制轨迹线，同时以不同颜色绘制速度、控制力、目标引导力、避障力、扰动力和合力向量，便于观察各类力对飞行轨迹的影响。", "轨迹与力向量可视化"),
        ("演示控制与状态面板", "HUD 实时显示飞行模式、速度、高度、目标距离和采集状态；演示控制面板支持切换模式、开关可视化、清除轨迹和调整参数；帮助面板提供快捷键和操作说明。", "演示控制与状态面板"),
    ]
    for i, (title, text, fig) in enumerate(steps, 1):
        p(doc, f"{i}. {title}：{text}")
        add_placeholder(doc, i, fig)

    p(doc, "通过以上步骤，用户可以完成从启动、飞行控制、目标交互到可视化查看和参数调节的完整使用流程。")

    doc.add_heading("三、软件总体设计", level=1)
    doc.add_heading("1. 软件特点", level=2)
    features = [
        "界面流程清晰：系统包含主菜单和核心仿真场景两个主要界面，主菜单提供“开始探索”和“退出”按钮作为统一的导航入口；仿真场景集成飞行控制、目标交互、可视化显示和相机跟随等功能。用户在场景间的切换由场景管理模块自动处理，整体操作路径简洁明确，无需额外配置即可开始使用。",
        "飞行行为实时生成：蜜蜂的运动轨迹并非预先录制的固定动画，而是每一帧由控制力（玩家输入或自动控制）、目标引导力、环境扰动力、障碍规避力、地面规避力和姿态调节算法在运行时共同解算生成。这种方式使蜜蜂能够对不同输入和场景变化做出即时响应，产生多样化的飞行表现。",
        "手动与自动模式结合：系统同时支持键盘手动飞行和自动目标巡航两种运行模式，用户可在运行时随时切换。手动模式适合交互体验和操作教学，自动模式适合行为演示和长时间稳定观察，两种模式共享同一套飞行控制框架，保证了行为规则的一致性。",
        "目标搜索与采集闭环：系统内置目标搜索与采集的完整行为链路——根据当前场景中的目标列表和检测半径筛选可访问目标，自动计算目标吸引力并驱动蜜蜂趋近；当蜜蜂进入采集触发范围后，进入悬停采集状态，完成后更新已采集记录并继续搜索下一目标，形成“搜索—趋近—悬停—采集—再搜索”的闭环流程。",
        "环境扰动表现自然：系统通过振翅扰动参数模拟昆虫飞行中翅膀振动带来的周期性微位移，同时叠加基于多层 Perlin 噪声的 Curl Noise 三维旋度力场，使蜜蜂飞行轨迹在保持宏观方向的同时呈现自然的细微波动和随机偏移，避免机械式的匀速直线运动，增强视觉真实感。",
        "避障与地面规避能力：系统在自动飞行过程中持续进行前向射线检测，根据障碍物的方位和距离计算回避力向量，使蜜蜂能够绕开场景中的树木、建筑等障碍物；同时通过地面高度监测和向上修正力，防止蜜蜂在自动巡航时贴地飞行或穿越地形，保证仿真过程的稳定性和可观察性。",
        "运行时可视化丰富：系统提供多维度的运行时可视化功能，包括飞行轨迹线（按时间间隔采样并以渐变色绘制）、速度向量、玩家控制力向量、目标引导力向量、障碍规避力向量、地面规避力向量、振翅扰动力向量、噪声扰动力向量以及合力向量。各力向量以不同颜色区分，便于观察和分析飞行过程中各类力对运动轨迹的贡献。",
        "演示辅助完整：系统提供运行时 HUD 状态面板，实时显示飞行模式、当前速度、飞行高度、目标距离、采集计数和操作快捷键提示；同时提供帮助面板说明各功能的使用方式，以及力学参数调节面板支持运行时调整飞行和显示参数。可视化开关和模式切换均配有界面按钮和快捷键双入口，方便在录屏、截图和现场演示场景下快速操作。",
        "模块化结构明确：系统将飞行控制、目标行为管理、轨迹记录、可视化开关控制、相机跟随、环境扰动计算和主菜单导航等功能拆分为独立的程序模块，各模块通过公开属性和组件引用进行协作，功能边界清晰。这种设计降低了模块间的耦合度，当需要修改某一功能（如调整飞行参数或增加新的可视化内容）时，只需关注对应模块，便于后续维护和功能扩展。",
    ]
    for item in features:
        bullet(doc, item)

    doc.add_heading("2. 功能模块设计", level=2)
    modules = [
        ["蜜蜂飞行控制模块", "BeeSimulation.cs", "处理手动/自动模式、速度积分、姿态俯仰侧倾、振翅扰动、噪声扰动、目标引导、避障和近地修正。"],
        ["目标搜索与采集模块", "BeeTargetController.cs", "管理目标列表、自动目标选择、目标悬停、玩家采集、已访问目标、稳定路线复现和采集计数。"],
        ["轨迹记录模块", "BeeSimulationManager.cs", "解析蜜蜂实例、创建 LineRenderer、按时间间隔记录轨迹点、清除轨迹、应用演示预设样式。"],
        ["可视化控制模块", "VisualizationManager.cs", "统一开启或关闭调试射线、运行时力向量、噪声显示和轨迹线显示。"],
        ["演示 HUD 模块", "BeePresentationHUD.cs", "提供状态文本、快捷键提示、导航按钮、帮助面板、力学参数面板、模式切换和可视化切换。"],
        ["相机跟随模块", "CameraFollow.cs", "根据目标、偏移模式和边界限制平滑跟随蜜蜂，支持固定旋转或朝向目标。"],
        ["环境扰动模块", "CurlNoiseField.cs", "基于多层 Perlin 噪声计算三维 Curl Noise 力场，为飞行轨迹提供动态扰动。"],
        ["主菜单控制模块", "MainMenuController.cs", "负责从首页加载 Demo 场景，以及在不同运行环境中执行退出逻辑。"],
    ]
    add_table(doc, ["模块名称", "对应脚本", "主要职责"], modules, widths=[3.3, 4.1, 8.5])

    doc.add_heading("3. 技术介绍", level=2)
    tech_sections = [
        ("1）Unity 三维场景与场景管理技术", "系统采用 Unity 引擎构建三维虚拟场景，使用 MainMenu.unity 承载软件首页和导航入口，使用 Demo.unity 承载蜜蜂飞行仿真主体。软件通过 UnityEngine.SceneManagement 进行场景切换，保证用户能够从主菜单进入仿真演示并在演示结束后返回。"),
        ("2）蜜蜂飞行动力合成技术", "BeeSimulation 将飞行过程抽象为多个力的合成，包括玩家输入控制力、目标引导力、地面规避力、障碍规避力、振翅扰动力、噪声扰动力和合力。系统在 FixedUpdate 中计算并限制总力，再更新刚体速度和位置，从而形成连续的飞行轨迹。"),
        ("3）姿态与振翅扰动表现技术", "系统根据飞行方向、速度、转向输入、爬升/下降状态和避障强度计算视觉模型的俯仰与侧倾，使蜜蜂在加速、转弯、上升、下降和避障时呈现更自然的姿态变化。振翅扰动参数用于模拟昆虫飞行中的细微周期波动。"),
        ("4）目标搜索、悬停与路线循环技术", "BeeTargetController 根据目标检测半径、候选过滤条件和已访问目标集合寻找当前目标；当蜜蜂接近目标时，系统可进入悬停状态，并围绕目标点生成小范围环绕或微悬停位置。自动模式下还可记录访问顺序并在下一轮稳定复现路线。"),
        ("5）避障与近地飞行控制技术", "系统通过前向射线检测识别障碍方向和距离，并根据障碍规避权重生成回避力；同时根据地面规避高度生成向上修正，避免自动飞行时穿越障碍或过低贴近地面。该机制提高了仿真过程的稳定性和可观察性。"),
        ("6）Curl Noise 环境扰动技术", "CurlNoiseField 通过多组 Perlin Noise 采样构建三维噪声梯度，并计算旋度型扰动力。噪声场随时间更新，使蜜蜂飞行轨迹具备非线性、微随机和自然波动效果，增强虚拟飞行的真实感。"),
        ("7）轨迹与力向量可视化技术", "系统使用 LineRenderer 绘制飞行轨迹和运行时向量。轨迹模块按照固定时间间隔采样蜜蜂位置并保留一定时长的历史点；力向量模块将速度、控制力、引导力、避障力、扰动力和合力以不同颜色显示，便于用户理解仿真计算过程。"),
        ("8）HUD 与运行时交互技术", "BeePresentationHUD 在运行时创建 Canvas、文本、按钮、帮助面板和参数滑块，提供快捷键 F1/F2/F3/Tab 等操作入口。用户可以在不离开仿真场景的情况下查看状态、切换模式、清理轨迹、打开帮助和调整演示参数。"),
        ("9）模块化软件架构", "系统将飞行控制、目标行为、轨迹显示、可视化开关、相机跟随和菜单导航拆分为独立脚本。各模块通过公开属性、组件引用和 Unity 生命周期方法协同工作，降低了功能耦合度，提高了后续扩展和维护的便利性。"),
    ]
    for title, text in tech_sections:
        doc.add_heading(title, level=3)
        p(doc, text)

    doc.add_heading("4. 主要数据与流程", level=2)
    flow_rows = [
        ["启动阶段", "加载 MainMenu 场景，显示标题、按钮和主菜单交互入口。"],
        ["进入阶段", "点击开始探索，加载 Demo 场景，解析蜜蜂实例、相机、目标点和管理脚本。"],
        ["输入阶段", "根据手动输入或自动模式生成控制指令与目标信息。"],
        ["计算阶段", "合成控制力、引导力、避障力、地面规避力、振翅扰动力和噪声扰动力。"],
        ["更新阶段", "更新刚体速度、位置、姿态、目标状态、采集计数和轨迹数据。"],
        ["显示阶段", "通过 HUD、轨迹线、力向量、射线和相机跟随展示运行状态。"],
        ["结束阶段", "用户返回主菜单或退出系统。"],
    ]
    add_table(doc, ["流程阶段", "处理内容"], flow_rows, widths=[3.2, 12.5])

    doc.add_heading("5. 面向领域和行业", level=2)
    domains = [
        "虚拟仿真教学：可用于生物行为、智能体控制、Unity 交互开发等课程中展示蜜蜂飞行、目标趋近和环境响应过程。",
        "科普展示与自然教育：可在课堂、展厅或科普活动中展示蜜蜂访花、飞行路径、悬停采集和环境扰动下的行为变化。",
        "智能体行为建模研究：可作为基于规则的自主飞行控制、目标搜索、路径循环和扰动响应模型的原型验证平台。",
        "数字内容与交互应用开发：可作为虚拟生态场景、三维昆虫飞行动画和交互式演示系统的实现参考。",
        "软件工程实践：系统脚本模块清晰，适合作为 Unity 项目中场景管理、运行时 UI、可视化调试和组件协作的示例。 ",
    ]
    for item in domains:
        bullet(doc, item)

    doc.add_heading("四、结论", level=1)
    for text in [
        "蜜蜂虚拟飞行行为仿真与可视化系统 V1.0 以 Unity 三维场景为载体，以 C# 脚本实现飞行控制、目标搜索、采集悬停、轨迹记录、环境扰动和运行时可视化等功能，形成了较完整的蜜蜂虚拟飞行行为仿真流程。",
        "系统能够在手动交互和自动巡航两种模式下运行，通过 HUD 状态面板、力向量、轨迹线和帮助界面增强演示解释性，符合软件著作权申报材料对软件功能、使用流程和技术实现说明的要求。",
    ]:
        p(doc, text)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
